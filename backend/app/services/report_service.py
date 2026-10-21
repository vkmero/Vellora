from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.exceptions import BadRequestError, NotFoundError
from app.models.analysis import Analysis
from app.models.approval import Approval
from app.models.cleaning_action import CleaningAction
from app.models.dataset import Dataset
from app.models.insight import Insight
from app.models.prediction import Prediction
from app.models.quality_issue import QualityIssue
from app.models.report import Report
from app.models.user import User
from app.models.validation import Validation
from app.services.file_service import EXCEL_DIR, PDF_DIR, ensure_storage
from app.services.notification_service import notify
from app.utils.helpers import utcnow


def generate_report(db: Session, dataset: Dataset, user_id: int, report_type: str, fmt: str) -> Report:
    ensure_storage()
    if report_type not in {"quality", "analytics", "complete"}:
        raise BadRequestError("Report type must be quality, analytics, or complete.")
    stamp = utcnow().strftime("%Y%m%d%H%M%S")
    if fmt == "pdf":
        path = PDF_DIR / f"{dataset.id}_{report_type}_{stamp}.pdf"
        _write_pdf(path, dataset, db, report_type, user_id)
    elif fmt in {"xlsx", "excel"}:
        path = EXCEL_DIR / f"{dataset.id}_{report_type}_{stamp}.xlsx"
        _write_excel(path, dataset, db, report_type)
    elif fmt == "csv":
        path = EXCEL_DIR / f"{dataset.id}_{report_type}_{stamp}.csv"
        _write_csv(path, dataset, db, report_type)
    else:
        raise BadRequestError("Format must be pdf, excel, or csv.")

    mgr_id = None
    if dataset.project and dataset.project.assigned_manager_id:
        mgr_id = dataset.project.assigned_manager_id
    else:
        active_mgr = db.query(User).filter(User.role.in_(["manager", "admin"]), User.is_active.is_(True)).first()
        if active_mgr:
            mgr_id = active_mgr.id

    record = Report(
        dataset_id=dataset.id,
        project_id=dataset.project_id,
        report_type=report_type,
        format=fmt,
        file_path=str(path),
        generated_by=user_id,
        delivered_to_manager_id=mgr_id,
        delivery_status="delivered",
        delivered_at=utcnow(),
        manager_review_status="pending_review",
    )
    db.add(record)
    db.flush()

    if mgr_id:
        notify(
            db,
            mgr_id,
            "New Deliverable Report Received",
            f"A {report_type.title()} ({fmt.upper()}) deliverable report was generated for dataset '{dataset.name}'.",
            "report_delivered",
        )

    return record


def serialize_report(record: Report, db: Session) -> dict:
    dataset = db.get(Dataset, record.dataset_id)
    generator = db.get(User, record.generated_by) if record.generated_by else None
    reviewer = db.get(User, record.delivered_to_manager_id) if record.delivered_to_manager_id else None
    return {
        "id": record.id,
        "project_id": record.project_id,
        "dataset_id": record.dataset_id,
        "report_type": record.report_type,
        "format": record.format or "pdf",
        "file_path": record.file_path,
        "generated_by": record.generated_by,
        "delivered_to_manager_id": record.delivered_to_manager_id,
        "delivery_status": record.delivery_status,
        "delivered_at": record.delivered_at,
        "delivery_error": record.delivery_error,
        "manager_review_status": record.manager_review_status,
        "manager_feedback": record.manager_feedback,
        "reviewed_at": record.reviewed_at,
        "created_at": record.created_at,
        "dataset_name": dataset.name if dataset else None,
        "generator_name": generator.name if generator else None,
        "reviewer_name": reviewer.name if reviewer else None,
    }


def review_report(db: Session, report_id: int, manager: User, review_status: str, feedback: str) -> Report:
    record = db.get(Report, report_id)
    if not record:
        raise NotFoundError("Report not found.")
    record.manager_review_status = review_status
    record.manager_feedback = feedback
    record.reviewed_at = utcnow()
    db.flush()

    if record.generated_by:
        notify(
            db,
            record.generated_by,
            f"Deliverable Review: {review_status.replace('_', ' ').title()}",
            f"Manager {manager.name} has reviewed your report: '{feedback}'",
            "report_reviewed",
        )
    return record


def _write_pdf(path: Path, dataset: Dataset, db: Session, report_type: str, user_id: int) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "VelloraTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        textColor=colors.HexColor("#0f172a"),
        alignment=0,
        spaceAfter=14,
    )
    h2_style = ParagraphStyle(
        "VelloraH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        textColor=colors.HexColor("#0369a1"),
        spaceBefore=12,
        spaceAfter=6,
    )
    normal_style = ParagraphStyle(
        "VelloraNormal",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
    )

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    story = [
        Paragraph("VELLORA ENTERPRISE DOSSIER", ParagraphStyle("Brand", fontName="Helvetica-Bold", fontSize=10, textColor=colors.HexColor("#0284c7"))),
        Paragraph(f"{report_type.title()} Governance & Analytics Report", title_style),
        Spacer(1, 8),
    ]

    # Executive Overview Table
    owner = db.get(User, dataset.owner_id)
    generator = db.get(User, user_id)
    overview_data = [
        ["Dataset Name", dataset.name, "Quality Index", f"{dataset.quality_score if dataset.quality_score is not None else 'N/A'}%"],
        ["Format / Size", f"{dataset.file_type.upper()} ({round(dataset.file_size / 1024, 1)} KB)", "Lifecycle State", dataset.status.replace("_", " ").title()],
        ["Rows / Columns", f"{dataset.row_count} rows, {dataset.column_count} columns", "Generated By", generator.name if generator else "Analyst"],
        ["Owner", owner.name if owner else "Analyst", "Generated At", utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")],
    ]
    t = Table(overview_data, colWidths=[100, 160, 100, 160])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0f172a")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    # Quality Section
    if report_type in {"quality", "complete"}:
        story.append(Paragraph("1. Data Quality & Issue Audit", h2_style))
        issues = db.query(QualityIssue).filter(QualityIssue.dataset_id == dataset.id).all()
        if not issues:
            story.append(Paragraph("No quality issues detected. All quality thresholds satisfied.", normal_style))
        else:
            issue_rows = [["Issue Type", "Column", "Severity", "Affected", "Description"]]
            for issue in issues[:15]:
                issue_rows.append([
                    issue.issue_type.replace("_", " ").title(),
                    issue.column_name or "Dataset",
                    issue.severity,
                    f"{issue.affected_rows} ({issue.percentage}%)",
                    issue.description[:60] + "..." if len(issue.description) > 60 else issue.description,
                ])
            it = Table(issue_rows, colWidths=[100, 80, 55, 65, 220])
            it.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(it)
        story.append(Spacer(1, 10))

    # Governance & Cleaning Section
    if report_type in {"complete"}:
        story.append(Paragraph("2. Transformation & Governance Trail", h2_style))
        actions = db.query(CleaningAction).filter(CleaningAction.dataset_id == dataset.id).all()
        if actions:
            act_rows = [["Action", "Column", "Before", "After", "Affected", "Description"]]
            for a in actions:
                act_rows.append([
                    a.action_type.replace("_", " ").title(),
                    a.column_name or "All",
                    str(a.before_count),
                    str(a.after_count),
                    str(a.affected_rows),
                    a.description[:55],
                ])
            at = Table(act_rows, colWidths=[90, 70, 45, 45, 45, 225])
            at.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(at)
        else:
            story.append(Paragraph("No automated cleaning operations applied prior to approval.", normal_style))
        story.append(Spacer(1, 10))

    # Insights Section
    if report_type in {"analytics", "complete"}:
        story.append(Paragraph("3. Algorithmic Business Insights", h2_style))
        insights = db.query(Insight).filter(Insight.dataset_id == dataset.id).all()
        if not insights:
            story.append(Paragraph("No business insights generated for this dataset.", normal_style))
        for insight in insights:
            story.append(Paragraph(f"&bull; <b>{insight.title}</b> ({insight.severity} severity): {insight.description}", normal_style))
            if insight.recommendation:
                story.append(Paragraph(f"&nbsp;&nbsp;<i>Recommendation: {insight.recommendation}</i>", ParagraphStyle("Rec", parent=normal_style, textColor=colors.HexColor("#0369a1"))))
        story.append(Spacer(1, 10))

    # Machine Learning Section
    if report_type in {"complete"}:
        prediction = db.query(Prediction).filter(Prediction.dataset_id == dataset.id).order_by(Prediction.created_at.desc()).first()
        if prediction:
            story.append(Paragraph("4. Predictive Machine Learning Dossier", h2_style))
            pred_lines = [
                f"Model: {prediction.model_type.replace('_', ' ').title()} | Target Column: {prediction.target_column}",
                f"Features: {', '.join(prediction.features or [])}",
                f"Metrics: {prediction.metrics}",
            ]
            for pl in pred_lines:
                story.append(Paragraph(pl, normal_style))

    doc.build(story)


def _write_excel(path: Path, dataset: Dataset, db: Session, report_type: str) -> None:
    wb = Workbook()
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    def style_header(sheet):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric / Attribute", "Value"])
    ws.append(["Dataset Name", dataset.name])
    ws.append(["Lifecycle State", dataset.status])
    ws.append(["Quality Score", dataset.quality_score])
    ws.append(["Row Count", dataset.row_count])
    ws.append(["Column Count", dataset.column_count])
    ws.append(["File Size Bytes", dataset.file_size])
    ws.append(["Generated Timestamp", utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")])
    style_header(ws)

    # Sheet 2: Quality Issues
    if report_type in {"quality", "complete"}:
        issues_ws = wb.create_sheet("Quality Issues")
        issues_ws.append(["Issue Type", "Column Name", "Severity", "Affected Rows", "Percentage", "Description", "Recommendation"])
        for issue in db.query(QualityIssue).filter(QualityIssue.dataset_id == dataset.id).all():
            issues_ws.append([issue.issue_type, issue.column_name, issue.severity, issue.affected_rows, issue.percentage, issue.description, issue.recommendation])
        style_header(issues_ws)

    # Sheet 3: Cleaning Transformations
    if report_type in {"complete"}:
        clean_ws = wb.create_sheet("Cleaning Transformations")
        clean_ws.append(["Action Type", "Column Name", "Before Count", "After Count", "Affected Rows", "Description"])
        for a in db.query(CleaningAction).filter(CleaningAction.dataset_id == dataset.id).all():
            clean_ws.append([a.action_type, a.column_name, a.before_count, a.after_count, a.affected_rows, a.description])
        style_header(clean_ws)

    # Sheet 4: Business Insights
    if report_type in {"analytics", "complete"}:
        insight_ws = wb.create_sheet("Business Insights")
        insight_ws.append(["Title", "Category", "Severity", "Description", "Recommendation"])
        for insight in db.query(Insight).filter(Insight.dataset_id == dataset.id).all():
            insight_ws.append([insight.title, insight.category, insight.severity, insight.description, insight.recommendation])
        style_header(insight_ws)

    # Sheet 5: Predictions
    if report_type in {"complete"}:
        preds = db.query(Prediction).filter(Prediction.dataset_id == dataset.id).all()
        if preds:
            pred_ws = wb.create_sheet("Predictions")
            pred_ws.append(["Model Type", "Target Column", "Features", "Metrics"])
            for p in preds:
                pred_ws.append([p.model_type, p.target_column, ", ".join(p.features or []), str(p.metrics)])
            style_header(pred_ws)

    wb.save(path)


def _write_csv(path: Path, dataset: Dataset, db: Session, report_type: str) -> None:
    issues = db.query(QualityIssue).filter(QualityIssue.dataset_id == dataset.id).all()
    lines = ["issue_type,column,severity,affected_rows,percentage,description"]
    for issue in issues:
        desc = (issue.description or "").replace(",", ";")
        lines.append(f"{issue.issue_type},{issue.column_name},{issue.severity},{issue.affected_rows},{issue.percentage},{desc}")
    path.write_text("\n".join(lines), encoding="utf-8")


def resolve_report_file(db: Session, report_id: int) -> Report:
    report = db.get(Report, report_id)
    if not report:
        raise NotFoundError("Report record was not found.")

    file_path = Path(report.file_path).resolve()
    base_dirs = [PDF_DIR.resolve(), EXCEL_DIR.resolve()]
    is_safe = any(file_path.is_relative_to(base) for base in base_dirs)
    if not is_safe or not file_path.exists():
        raise NotFoundError("Report file was not found or access is restricted.")
    return report
